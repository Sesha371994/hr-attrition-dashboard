"""
HR Attrition Analytics — full pipeline (load -> SQL clean -> Python
transform -> metrics -> Excel dashboard) in one script for a compact
portfolio deliverable.
"""
import sqlite3
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.utils import get_column_letter

# ---------------- LOAD ----------------
conn = sqlite3.connect("hr.db")
raw = pd.read_csv("raw_data/hr_employees.csv")
raw.to_sql("raw_hr", conn, if_exists="replace", index=False)

with open("sql/cleaning.sql") as f:
    conn.executescript(f.read())
conn.commit()

clean = pd.read_sql("SELECT * FROM clean_hr", conn)
print(f"Raw rows: {len(raw)} -> Clean rows: {len(clean)}")

# ---------------- PYTHON TRANSFORM (dates) ----------------
def parse_date(s):
    if pd.isna(s) or s == "":
        return pd.NaT
    for fmt in ["%Y-%m-%d", "%d/%m/%Y", "%m-%d-%Y"]:
        try:
            return pd.to_datetime(s, format=fmt)
        except (ValueError, TypeError):
            continue
    return pd.NaT

clean["join_date"] = clean["join_date_raw"].apply(parse_date)
clean["exit_date"] = clean["exit_date_raw"].apply(parse_date)
clean = clean.drop(columns=["join_date_raw", "exit_date_raw"])

# tenure in months (for attrited: join->exit, for active: join->today)
today = pd.Timestamp("2025-09-08")
clean["end_ref"] = clean["exit_date"].fillna(today)
clean["tenure_months"] = ((clean["end_ref"] - clean["join_date"]).dt.days / 30.44).round(1)
clean = clean.drop(columns=["end_ref"])

clean.to_csv("exports/hr_clean.csv", index=False)

# ---------------- METRICS ----------------
total_emp = len(clean)
attrited_count = (clean["attrited"] == "Yes").sum()
attrition_rate = attrited_count / total_emp
avg_tenure = clean["tenure_months"].mean()
avg_satisfaction = clean["satisfaction_score"].mean()

dept_attrition = clean.groupby("department").apply(
    lambda d: (d["attrited"] == "Yes").sum() / len(d)
).round(3).sort_values(ascending=False)

print(f"Attrition Rate: {attrition_rate:.1%}")
print(f"Avg Tenure: {avg_tenure:.1f} months")
print("\nAttrition by Dept:\n", dept_attrition)

# ---------------- EXCEL DASHBOARD ----------------
wb = Workbook()
FONT = "Arial"
HEADER_FILL = PatternFill(start_color="1E3C72", end_color="1E3C72", fill_type="solid")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name=FONT, bold=True, size=14, color="1E3C72")
LABEL_FONT = Font(name=FONT, bold=True, size=11)
NORMAL_FONT = Font(name=FONT, size=10)

ws_data = wb.active
ws_data.title = "Raw_Data"
cols = ["employee_id", "name", "gender", "age", "department", "education",
        "monthly_salary", "satisfaction_score", "attrited", "tenure_months"]
for c, col in enumerate(cols, 1):
    cell = ws_data.cell(1, c, col.replace("_", " ").title())
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
for r, row in enumerate(clean[cols].itertuples(index=False), 2):
    for c, val in enumerate(row, 1):
        ws_data.cell(r, c, val).font = NORMAL_FONT
for c, col in enumerate(cols, 1):
    ws_data.column_dimensions[get_column_letter(c)].width = max(14, len(col) + 4)
n_rows = len(clean) + 1
ws_data.freeze_panes = "A2"

ws = wb.create_sheet("Dashboard")
ws["B2"] = "HR Attrition Analytics Dashboard"
ws["B2"].font = TITLE_FONT

dept_col = cols.index("department") + 1
attr_col = cols.index("attrited") + 1
sat_col = cols.index("satisfaction_score") + 1
sal_col = cols.index("monthly_salary") + 1
dept_l = get_column_letter(dept_col)
attr_l = get_column_letter(attr_col)
sat_l = get_column_letter(sat_col)
sal_l = get_column_letter(sal_col)

r_dept = f"Raw_Data!${dept_l}$2:${dept_l}${n_rows}"
r_attr = f"Raw_Data!${attr_l}$2:${attr_l}${n_rows}"
r_sat = f"Raw_Data!${sat_l}$2:${sat_l}${n_rows}"
r_sal = f"Raw_Data!${sal_l}$2:${sal_l}${n_rows}"

ws["B6"] = "Total Employees"
ws["D6"] = "Attrition Rate"
ws["F6"] = "Avg Satisfaction"
ws["H6"] = "Avg Salary"
for c in ["B6", "D6", "F6", "H6"]:
    ws[c].font = LABEL_FONT

ws["B7"] = f'=COUNTA({r_attr})'
ws["D7"] = f'=COUNTIF({r_attr},"Yes")/B7'
ws["D7"].number_format = "0.0%"
ws["F7"] = f'=AVERAGEIF({r_sat},"<>",{r_sat})'
ws["F7"].number_format = "0.00"
ws["H7"] = f'=AVERAGEIF({r_sal},"<>",{r_sal})'
ws["H7"].number_format = '"₹"#,##0'
for c in ["B7", "D7", "F7", "H7"]:
    ws[c].font = Font(name=FONT, size=16, bold=True, color="1E3C72")

ws["B10"] = "Attrition Rate by Department"
ws["B10"].font = LABEL_FONT
ws["B11"] = "Department"
ws["C11"] = "Attrition Rate"
ws["B11"].font = HEADER_FONT
ws["C11"].font = HEADER_FONT
ws["B11"].fill = HEADER_FILL
ws["C11"].fill = HEADER_FILL

depts = sorted(clean["department"].unique().tolist())
for i, d in enumerate(depts, 12):
    ws[f"B{i}"] = d
    ws[f"C{i}"] = f'=COUNTIFS({r_dept},B{i},{r_attr},"Yes")/COUNTIF({r_dept},B{i})'
    ws[f"C{i}"].number_format = "0.0%"
last_dept_row = 11 + len(depts)

ws["E10"] = "Headcount by Department"
ws["E10"].font = LABEL_FONT
ws["E11"] = "Department"
ws["F11"] = "Count"
ws["E11"].font = HEADER_FONT
ws["F11"].font = HEADER_FONT
ws["E11"].fill = HEADER_FILL
ws["F11"].fill = HEADER_FILL
for i, d in enumerate(depts, 12):
    ws[f"E{i}"] = d
    ws[f"F{i}"] = f'=COUNTIF({r_dept},E{i})'
last_hc_row = 11 + len(depts)

bar = BarChart()
bar.title = "Attrition Rate by Department"
data_ref = Reference(ws, min_col=3, min_row=11, max_row=last_dept_row)
cats_ref = Reference(ws, min_col=2, min_row=12, max_row=last_dept_row)
bar.add_data(data_ref, titles_from_data=True)
bar.set_categories(cats_ref)
bar.width = 14
bar.height = 8
ws.add_chart(bar, "B18")

pie = PieChart()
pie.title = "Headcount by Department"
data_ref2 = Reference(ws, min_col=6, min_row=11, max_row=last_hc_row)
cats_ref2 = Reference(ws, min_col=5, min_row=12, max_row=last_hc_row)
pie.add_data(data_ref2, titles_from_data=True)
pie.set_categories(cats_ref2)
pie.width = 14
pie.height = 8
ws.add_chart(pie, "F18")

for col, w in [("B", 16), ("C", 16), ("D", 14), ("E", 16), ("F", 14), ("G", 14), ("H", 14)]:
    ws.column_dimensions[col].width = w

wb.save("exports/HR_Attrition_Dashboard.xlsx")
print("\nSaved: exports/HR_Attrition_Dashboard.xlsx")
conn.close()
